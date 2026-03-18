from __future__ import annotations

import csv
import sys
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from typing import Iterable

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageTk

from vote_logic import VotingEngine
from wjx_parser import parse_wjx_ranking_text


class VotingApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("雷达院奖学金专家/学生投票系统")
        self.root.geometry("1180x820")
        self.root.configure(bg="#e9f5ee")

        self.candidates: list[str] = []
        self.engine: VotingEngine | None = None
        self.voter_weights: dict[str, float] = {}

        self._build_ui()

    def _resource_path(self, relative: str) -> Path:
        """支持源码运行与 PyInstaller onefile 运行时资源路径。"""
        base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
        return base / relative

    def _create_gradient_photo(self, width: int, height: int) -> ImageTk.PhotoImage:
        start = (13, 109, 61)   # BIT 绿色
        end = (170, 41, 37)     # BIT 校徽红
        img = Image.new("RGB", (max(1, width), max(1, height)), start)
        draw = ImageDraw.Draw(img)
        for y in range(max(1, height)):
            t = y / max(height - 1, 1)
            r = int(start[0] * (1 - t) + end[0] * t)
            g = int(start[1] * (1 - t) + end[1] * t)
            b = int(start[2] * (1 - t) + end[2] * t)
            draw.line([(0, y), (max(1, width), y)], fill=(r, g, b))
        return ImageTk.PhotoImage(img)

    def _load_bit_logo(self, size: int = 72) -> ImageTk.PhotoImage:
        logo_path = self._resource_path("assets/bit_logo.png")
        if logo_path.exists():
            img = Image.open(logo_path).convert("RGBA").resize((size, size), Image.Resampling.LANCZOS)
            return ImageTk.PhotoImage(img)

        # fallback: 生成简化占位徽标
        img = Image.new("RGBA", (size, size), (0, 0, 0, 0))
        draw = ImageDraw.Draw(img)
        draw.ellipse((2, 2, size - 2, size - 2), fill=(170, 41, 37, 255), outline=(238, 202, 181, 255), width=4)
        draw.ellipse((12, 12, size - 12, size - 12), fill=(13, 109, 61, 255))
        draw.text((size // 2 - 12, size // 2 - 8), "BIT", fill=(255, 255, 255, 255))
        return ImageTk.PhotoImage(img)

    def _refresh_banner(self) -> None:
        width = max(200, self.banner_canvas.winfo_width())
        height = self.banner_height
        self.banner_bg = self._create_gradient_photo(width, height)
        self.banner_canvas.delete("all")
        self.banner_canvas.create_image(0, 0, image=self.banner_bg, anchor="nw")
        self.banner_canvas.create_image(20, height // 2, image=self.bit_logo, anchor="w")
        self.banner_canvas.create_text(
            110,
            height // 2,
            text="雷达院奖学金专家/学生投票系统（问卷星联动）",
            font=("Microsoft YaHei", 18, "bold"),
            fill="#ffffff",
            anchor="w",
        )

    def _on_root_resize(self, event: tk.Event) -> None:
        if event.widget is self.root:
            self._refresh_banner()


    def _build_ui(self) -> None:
        self.banner_height = 96
        self.banner_canvas = tk.Canvas(self.root, height=self.banner_height, highlightthickness=0, bd=0)
        self.banner_canvas.pack(fill="x", padx=12, pady=(10, 6))
        self.bit_logo = self._load_bit_logo(72)
        self.root.bind("<Configure>", self._on_root_resize)
        self.root.after(50, self._refresh_banner)

        weight_frame = ttk.LabelFrame(self.root, text="初评（权重设置 + 专家文件导入）")
        weight_frame.pack(fill="x", padx=12, pady=8)
        self.voter_name_var = tk.StringVar(value="")
        self.weight_var = tk.StringVar(value="1")
        ttk.Label(weight_frame, text="投票人姓名：").pack(side="left", padx=6)
        ttk.Entry(weight_frame, textvariable=self.voter_name_var, width=18).pack(side="left", padx=6)
        ttk.Label(weight_frame, text="权重：").pack(side="left", padx=6)
        ttk.Entry(weight_frame, textvariable=self.weight_var, width=8).pack(side="left", padx=6)
        ttk.Button(weight_frame, text="新增/更新权重", command=self.set_voter_weight).pack(side="left", padx=8)
        ttk.Button(weight_frame, text="重置权重", command=self.reset_weights).pack(side="left", padx=8)
        ttk.Button(weight_frame, text="Top-X次数统计", command=self.count_topx_frequency).pack(side="left", padx=8)
        ttk.Button(weight_frame, text="导入专家文件", command=self.import_expert_file).pack(side="left", padx=8)
        ttk.Button(weight_frame, text="重置数据", command=self.reset_all).pack(side="left", padx=8)

        import_frame = ttk.LabelFrame(self.root, text="终评（问卷星结果导入与结算）")
        import_frame.pack(fill="x", padx=12, pady=8)

        ttk.Button(import_frame, text="导入专家文件", command=self.import_expert_file).pack(side="left", padx=6, pady=8)
        ttk.Button(import_frame, text="导入学生文件", command=self.import_student_file).pack(side="left", padx=6, pady=8)
        ttk.Button(import_frame, text="重置数据", command=self.reset_all).pack(side="left", padx=6, pady=8)
        ttk.Button(import_frame, text="专家票结算", command=self.settle_expert).pack(side="left", padx=6, pady=8)
        ttk.Button(import_frame, text="学生票结算并计算最终排名", command=self.settle_final).pack(side="left", padx=6, pady=8)

        self.status_var = tk.StringVar(value="状态：请导入问卷星结果并开始结算。")
        ttk.Label(self.root, textvariable=self.status_var).pack(fill="x", padx=12)

        self.output = tk.Text(self.root, height=24, font=("Consolas", 11))
        self.output.pack(fill="both", expand=True, padx=12, pady=10)
        self.output.insert(
            "end",
            "欢迎使用系统。\n"
            "导入格式：问卷星固定结构，第7列是排序结果，第8列是“2、您的姓名”。\n",
        )

    def set_voter_weight(self) -> None:
        name = self.voter_name_var.get().strip()
        if not name:
            messagebox.showwarning("提示", "请输入投票人姓名")
            return
        try:
            weight = float(self.weight_var.get().strip())
        except ValueError:
            messagebox.showwarning("提示", "权重必须是数字")
            return
        if weight <= 0:
            messagebox.showwarning("提示", "权重必须大于0")
            return

        self.voter_weights[name] = weight
        self.status_var.set("状态：权重设置成功。")
        self.output.insert("end", f"已设置权重：{name} -> {weight}\n")


    def reset_weights(self) -> None:
        self.voter_weights.clear()
        self.voter_name_var.set("")
        self.weight_var.set("1")
        self.status_var.set("状态：已重置所有投票人权重。")
        self.output.insert("end", "已清空全部投票人权重设置。\n")


    def count_topx_frequency(self) -> None:
        if self.engine is None:
            messagebox.showwarning("提示", "请先导入数据")
            return

        self.output.insert("end", "\n=== Top-X 选择次数统计（不算名次分）===\n")

        if self.engine.expert_ballots:
            counts, rows = self.engine.settle_expert_selection_count()
            self.output.insert("end", "\n专家票 Top-X 次数\n")
            self.output.insert("end", "排名\t候选人\t被选择次数\n")
            for r, name, c in rows:
                display = int(c) if abs(c - int(c)) < 1e-9 else round(c, 2)
                self.output.insert("end", f"{r}\t{name}\t{display}\n")
            self.output.insert("end", f"专家原始次数：{counts}\n")

        if self.engine.student_ballots:
            counts, rows = self.engine.settle_student_selection_count()
            self.output.insert("end", "\n学生票 Top-X 次数\n")
            self.output.insert("end", "排名\t候选人\t被选择次数\n")
            for r, name, c in rows:
                display = int(c) if abs(c - int(c)) < 1e-9 else round(c, 2)
                self.output.insert("end", f"{r}\t{name}\t{display}\n")
            self.output.insert("end", f"学生原始次数：{counts}\n")

        if not self.engine.expert_ballots and not self.engine.student_ballots:
            messagebox.showwarning("提示", "尚未导入任何票")
            return

        self.status_var.set("状态：已生成Top-X次数统计。")

    def _read_csv_rows(self, path: str) -> list[list[str]]:
        for enc in ["utf-8-sig", "gbk", "gb18030", "utf-8"]:
            try:
                with open(path, "r", encoding=enc, newline="") as f:
                    return list(csv.reader(f))
            except UnicodeDecodeError:
                continue
        raise ValueError("CSV编码无法识别，请尝试另存为 UTF-8 或 GBK")

    def _read_xlsx_rows(self, path: str) -> list[list[str]]:
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else str(v) for v in row])
        return rows

    def _iter_data_rows(self, path: str) -> Iterable[list[str]]:
        ext = Path(path).suffix.lower()
        if ext == ".xlsx":
            rows = self._read_xlsx_rows(path)
        else:
            rows = self._read_csv_rows(path)
        if not rows:
            return []
        return rows[1:]

    def _load_ballots_from_file(self, path: str) -> tuple[list[tuple[dict[str, int], str, float]], list[str], int]:
        raw_ballots: list[tuple[dict[str, int], str, float]] = []
        file_candidates_ordered: list[str] = []
        file_candidates_set: set[str] = set()
        skipped_rows = 0

        for row in self._iter_data_rows(path):
            if len(row) < 7:
                continue
            ranking_text = (row[6] or "").strip()
            voter_name = (row[7] or "").strip() if len(row) > 7 else ""
            if not ranking_text:
                continue

            ordered_names = parse_wjx_ranking_text(ranking_text)
            if not ordered_names or len(set(ordered_names)) != len(ordered_names):
                skipped_rows += 1
                continue

            for name in ordered_names:
                if name not in file_candidates_set:
                    file_candidates_set.add(name)
                    file_candidates_ordered.append(name)

            ballot = {name: rank for rank, name in enumerate(ordered_names, start=1)}
            weight = self.voter_weights.get(voter_name, 1.0)
            raw_ballots.append((ballot, voter_name, weight))

        if not raw_ballots:
            raise ValueError("文件中没有有效投票数据")

        return raw_ballots, file_candidates_ordered, skipped_rows

    def _import_file(self, role: str) -> None:
        path = filedialog.askopenfilename(
            title=f"选择{role}文件",
            filetypes=[("Excel/CSV Files", "*.xlsx *.csv"), ("All Files", "*.*")],
        )
        if not path:
            return

        try:
            rows, file_candidates, skipped_rows = self._load_ballots_from_file(path)
            if self.engine is None:
                self.candidates = file_candidates
                self.engine = VotingEngine(self.candidates)
                self.output.insert("end", f"自动识别候选人：{self.candidates}\n")
            else:
                added = self.engine.ensure_candidates(file_candidates)
                if added > 0:
                    self.candidates = self.engine.candidates
                    self.output.insert("end", f"检测到新增候选人 {added} 位，已自动扩展候选人池。\n")

            assert self.engine is not None
            if role == "专家":
                self.engine.expert_ballots = []
                self.output.insert("end", "已清空此前专家导入记录。\n")
                for ballot, voter_name, weight in rows:
                    self.engine.add_expert_ballot(ballot, weight=weight)
                    if weight != 1.0:
                        self.output.insert("end", f"专家票加权：{voter_name} 权重 {weight}\n")
                self.output.insert("end", f"已导入专家文件：{Path(path).name}，当前 {len(rows)} 票。\n")
            else:
                self.engine.student_ballots = []
                self.output.insert("end", "已清空此前学生导入记录。\n")
                for ballot, voter_name, weight in rows:
                    self.engine.add_student_ballot(ballot, weight=weight)
                    if weight != 1.0:
                        self.output.insert("end", f"学生票加权：{voter_name} 权重 {weight}\n")
                self.output.insert("end", f"已导入学生文件：{Path(path).name}，当前 {len(rows)} 票。\n")
            if skipped_rows > 0:
                self.output.insert("end", f"警告：已跳过 {skipped_rows} 条无效排序记录。\n")
            self.status_var.set("状态：文件导入成功。")
        except Exception as e:
            messagebox.showerror("导入失败", str(e))

    def import_expert_file(self) -> None:
        self._import_file("专家")

    def import_student_file(self) -> None:
        self._import_file("学生")

    def reset_all(self) -> None:
        self.engine = VotingEngine(self.candidates) if self.candidates else None
        self.output.delete("1.0", "end")
        self.output.insert("end", "已重置票数数据。\n")
        self.status_var.set("状态：票数已重置。")

    def _render_rank(self, title: str, rank_rows: list[tuple[int, str, float]]) -> None:
        self.output.insert("end", f"\n{title}\n")
        self.output.insert("end", "排名\t候选人\t总分(越低越好)\n")
        for r, name, score in rank_rows:
            self.output.insert("end", f"{r}\t{name}\t{score:.2f}\n")

    def settle_expert(self) -> None:
        if self.engine is None or not self.engine.expert_ballots:
            messagebox.showwarning("提示", "尚未导入专家票")
            return
        totals, rank_rows = self.engine.settle_experts()
        self._render_rank("专家票结算", rank_rows)
        self.status_var.set("状态：专家票已结算。")
        self.output.insert("end", f"专家票总分：{totals}\n")

    def settle_final(self) -> None:
        if self.engine is None or not self.engine.student_ballots:
            messagebox.showwarning("提示", "尚未导入学生票")
            return

        result = self.engine.final_result()
        self._render_rank("学生票结算", result["student_rank"])

        self.output.insert("end", "\n学生排名对专家总分的加权（减分）\n")
        for rank, name, _ in result["student_rank"]:
            delta = result["student_adjustment"][name]
            self.output.insert("end", f"学生第{rank}名：{name}，专家总分减{delta}\n")

        self._render_rank("最终排名（按调整后的专家总分）", result["final_rank"])
        self.status_var.set("状态：最终结果已生成。")


def main() -> None:
    root = tk.Tk()
    VotingApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
